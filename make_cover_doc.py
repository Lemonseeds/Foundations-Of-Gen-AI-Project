"""
Generates the submission Word document for Step 3 of the assignment:
YourName_RegNum_FGAI_Assignment2.docx

Usage:
    python make_cover_doc.py --name "Your Name" --reg "Your RegNum" --repo "https://github.com/YourUsername/YourRepoName"
"""
import argparse
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def build_submission_doc(
    student_name: str = "Ayrton",
    reg_number: str = "REG_NUMBER_HERE",
    github_repo_url: str = "https://github.com/Ayrton/ICT4403-FGAI-Assignment2-Autoregressive-Text-Generation",
    output_filename: str | None = None,
):
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_course = p_title.add_run("FOUNDATIONS OF GENERATIVE AI (ICT4403)\n")
    run_course.font.size = Pt(14)
    run_course.font.bold = True
    run_course.font.color.rgb = RGBColor(38, 70, 83)

    run_sub = p_title.add_run("Assignment 2: Autoregressive Text Generation\nSubmission Cover Sheet\n")
    run_sub.font.size = Pt(18)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(42, 157, 143)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(20)
    run_line = p_div.add_run("―" * 45)
    run_line.font.color.rgb = RGBColor(200, 200, 200)

    # Submission Details Table
    details = [
        ("Student Name", student_name),
        ("Registration Number", reg_number),
        ("Assignment Title", "Assignment 2: Comparative Autoregressive Text Generation using Recurrent Neural Networks"),
        (
            "Title of Developed System",
            "Autoregressive Next-Word Prediction & Multi-Temperature Generation System using Simple RNN, Bidirectional RNN, LSTM, and GRU Models",
        ),
        (
            "Name of Dataset Used",
            "The Devil's Dictionary (Ambrose Bierce, 1911) — Project Gutenberg eBook #972",
        ),
        (
            "Dataset Source Link",
            "https://www.gutenberg.org/ebooks/972 (Raw UTF-8: https://www.gutenberg.org/cache/epub/972/pg972.txt)",
        ),
        ("GitHub Repository Link", github_repo_url),
    ]

    table = doc.add_table(rows=len(details), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for i, (label, val) in enumerate(details):
        row = table.rows[i]
        
        # Label cell
        cell_lbl = row.cells[0]
        cell_lbl.width = Inches(2.3)
        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(6)
        p_lbl.paragraph_format.space_after = Pt(6)
        run_l = p_lbl.add_run(label)
        run_l.font.bold = True
        run_l.font.size = Pt(11)
        run_l.font.color.rgb = RGBColor(38, 70, 83)

        # Value cell
        cell_val = row.cells[1]
        cell_val.width = Inches(4.2)
        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_before = Pt(6)
        p_val.paragraph_format.space_after = Pt(6)
        run_v = p_val.add_run(val)
        run_v.font.size = Pt(11)
        if "http" in val:
            run_v.font.color.rgb = RGBColor(33, 102, 172)
            run_v.font.underline = True

    # Brief Abstract / Executive Summary
    p_sec = doc.add_paragraph()
    p_sec.paragraph_format.space_before = Pt(24)
    p_sec.paragraph_format.space_after = Pt(6)
    r_sec = p_sec.add_run("Executive Summary of Developed System & Findings")
    r_sec.font.bold = True
    r_sec.font.size = Pt(13)
    r_sec.font.color.rgb = RGBColor(38, 70, 83)

    summary_text = (
        "This project implements an end-to-end autoregressive language modeling system comparing four recurrent neural "
        "network architectures: Simple RNN, Bidirectional RNN (causally constrained to the context window), LSTM, and GRU. "
        "The models were evaluated on the complete cleaned corpus of Ambrose Bierce's 'The Devil's Dictionary' (Project Gutenberg #972), "
        "an atypical literary dataset featuring dual registers of cynical prose definitions and mock-heroic verse.\n\n"
        "Under identical experimental conditions (shared vocabulary of 4,731 words, context length 20, embedding dim 128, hidden dim 256, "
        "Adam optimizer, gradient clipping 1.0, and early stopping), the models achieved the following results:\n"
        "• LSTM achieved the lowest test loss (5.7027) and best test perplexity (299.67), successfully modeling long-range syntactic patterns.\n"
        "• GRU achieved the highest test top-1 prediction accuracy (18.11%) and top-5 accuracy (35.38%) while training ~27% faster than LSTM.\n"
        "• BiRNN attained a test perplexity of 373.67, outperforming Simple RNN (424.66) by leveraging forward and backward sweeps over the context window.\n"
        "• Temperature sampling analysis across T in {0.2, 0.5, 1.0, 1.5} demonstrated that T=0.2 induces phrase-level cyclic degeneration, "
        "T=0.5 produces optimal, stylistically coherent Biercian definitions, T=1.0 exhibits natural creative diversity (distinct-2 > 0.91), "
        "and T=1.5 causes excessive semantic divergence.\n\n"
        "All code, Jupyter notebook (FGAI_Assignment2.ipynb), saved model checkpoints, generated samples, and visual plots are organized "
        "in the accompanying GitHub repository."
    )

    p_sum = doc.add_paragraph(summary_text)
    p_sum.paragraph_format.space_after = Pt(12)
    p_sum.paragraph_format.line_spacing = 1.15

    # Determine filename
    safe_name = student_name.replace(" ", "_")
    safe_reg = reg_number.replace(" ", "_")
    if not output_filename:
        output_filename = f"{safe_name}_{safe_reg}_FGAI_Assignment2.docx"

    doc.save(output_filename)
    print(f"Generated Word document: {output_filename}")
    return output_filename


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--name", default="Ayrton", help="Student Name")
    parser.add_argument("--reg", default="REG_NUM", help="Registration Number")
    parser.add_argument("--repo", default="https://github.com/Ayrton/ICT4403-FGAI-Assignment2", help="GitHub repo link")
    parser.add_argument("--out", default=None, help="Output docx filename")
    args = parser.parse_args()

    build_submission_doc(args.name, args.reg, args.repo, args.out)
